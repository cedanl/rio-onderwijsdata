"""CREBO-lijst van SBB: genormaliseerde tabel, editie-cache en deterministische lookup.

Bouwt voort op ``sbb.resources("crebolijst")``; de bestaande ``sbb.load()`` blijft werken.

    from riodata import crebo

    crebo.ververs("codelijst_2025_april")      # download, valideer en cache (netwerk)
    r = crebo.zoek("25604")                    # 'latest' = nieuwste gevalideerde editie in de cache
    r["status"]                                # "found" | "not_found" | "ambiguous"

Vereist ``openpyxl`` (``pip install 'riodata[sbb]'``).

Ontwerpkeuzes:
- CREBO-codes blijven strings. In het echte bestand zijn ze deels getal en deels tekst met
  niet-brekende spaties; die worden schoongemaakt. Voorloopnullen die in een numerieke cel al
  verloren zijn, worden niet geraden of aangevuld.
- Ontbrekende waarden worden ``None``, rijen blijven behouden. Dubbele codes (over cohorten)
  blijven bestaan; de lookup geeft ze terug in plaats van de eerste.
- Bronkolommen blijven in ``bron`` bewaard. Herkent de loader de kopregel of het werkblad
  "Complete lijst" niet, dan volgt een ``CreboSchemaFout`` met wat er wel staat, geen gok.
- Geldigheid: de editie geldt vanaf de datum in de titelregel; einddata bestaan alleen per dossier
  (werkblad "Vervallen") en staan dan in ``geldig_tot``.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

# Het echte SBB-bestand (gecontroleerd op de edities april 2025, oktober 2025 en april 2026, plus
# 2021): werkblad "Complete lijst" (in 2021 niet het eerste werkblad); de titelregel
# ("Overzicht vastgestelde kwalificatiedossiers en kwalificaties geldig vanaf 01-08-2026") staat in
# dezelfde rij als de kolomkoppen Opleidingscode (2021: Crebonummer), Kwalificatie, Niveau,
# Prijsfactor, Soort opleiding, Beroepsvereisten, Leerweg. Daaronder staat een tweede kopregel voor
# het dossier (Opleidingscode, Prijsfactor, Kwalificatiedossier) en rijen voor opleidingsdomeinen.
# Codes zijn deels getal, deels tekst met niet-brekende spaties ('25950\xa0'). Er is geen begin- of
# einddatum per kwalificatie: de geldigheid is die van de editie (titelregel); einddata staan alleen
# per dossier in het werkblad "Vervallen".
WERKBLAD = "complete lijst"
VERVALLEN = "vervallen"
CODE_KOPPEN = ("opleidingscode", "crebonummer", "crebocode")
KOLOMKOPPEN = {
    "naam": ("kwalificatie",),
    "niveau": ("niveau",),
    "prijsfactor": ("prijsfactor",),
    "soort_opleiding": ("soortopleiding",),
    "beroepsvereisten": ("beroepsvereisten",),
    "leerweg": ("leerweg",),
}
VERPLICHT = ("crebo", "naam")
MIN_RIJEN = 100  # een CREBO-lijst met minder rijen is een afgekapt of verkeerd bestand
XLSX_CONTENTTYPES = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",  # zo levert kwalificatie-mijn.s-bb.nl sommige edities uit
)


class CreboFout(Exception):
    """Basis voor alle CREBO-loaderfouten."""


class CreboSchemaFout(CreboFout):
    """Het bestand heeft niet de verwachte kolommen of omvang."""


class CreboEditieOnbekend(CreboFout):
    """Geen (gevalideerde) editie beschikbaar."""


@dataclass(frozen=True)
class CreboLijst:
    editie: str
    rijen: tuple[dict, ...]
    kolommen: tuple[str, ...]
    bron_url: str | None = None
    sha256: str | None = None
    gecontroleerd_op: str | None = None
    geldig_vanaf: str | None = None
    vervallen: dict = field(default_factory=dict)


def _norm(kop) -> str:
    return re.sub(r"[\s\u202f\xa0_\-./]", "", str(kop or "").lower())


def _tekst(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # 25604.0 uit een numerieke cel
    s = str(v).replace("\xa0", " ").replace("\u202f", " ").strip()
    return s or None


def _getal(v) -> float | None:
    t = _tekst(v)
    try:
        return float(t.replace(",", ".")) if t else None
    except ValueError:
        return None


def _is_code(v) -> bool:
    t = _tekst(v)
    return bool(t) and t.isdigit()


def _datum(v) -> str | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = _tekst(v) or ""
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return s  # niet te lezen: ruwe tekst behouden, niet weggooien


def _werkblad(wb, naam: str):
    for ws in wb.worksheets:
        if ws.title.strip().lower() == naam:
            return ws
    return None


def _rijen(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _kopregel(rijen: list[tuple]) -> tuple[int, dict[str, int]]:
    """Vind de kopregel: de rij met een codekop plus Kwalificatie en Niveau."""
    for i, rij in enumerate(rijen[:15]):
        pos: dict[str, int] = {}
        for j, c in enumerate(rij):
            n = _norm(c)
            if n in CODE_KOPPEN and "crebo" not in pos:
                pos["crebo"] = j
            for doel, aliassen in KOLOMKOPPEN.items():
                if n in aliassen and doel not in pos:
                    pos[doel] = j
        if all(k in pos for k in VERPLICHT) and "niveau" in pos:
            return i, pos
    voorbeeld = [[_tekst(c) for c in r[:10]] for r in rijen[:4]]
    raise CreboSchemaFout(
        "Geen kopregel met Opleidingscode/Crebonummer, Kwalificatie en Niveau gevonden. "
        f"Eerste rijen: {voorbeeld}"
    )


def _dossierkolommen(rijen: list[tuple], na: int) -> dict[str, int]:
    """Tweede kopregel (Opleidingscode, Prijsfactor, Kwalificatiedossier) voor de dossierkolommen."""
    for rij in rijen[na + 1: na + 4]:
        pos = {}
        for j, c in enumerate(rij):
            n = _norm(c)
            if n in CODE_KOPPEN and "code" not in pos:
                pos["code"] = j
            elif n == "kwalificatiedossier":
                pos["naam"] = j
        if "naam" in pos:
            return pos
    return {}


def _vervallen(wb) -> dict[str, dict]:
    ws = _werkblad(wb, VERVALLEN)
    if ws is None:
        return {}
    rijen = _rijen(ws)
    for i, rij in enumerate(rijen[:10]):
        koppen = [_norm(c) for c in rij]
        if any(k.startswith("opleidingscode") or k.startswith("crebo") for k in koppen) and "dossiernaam" in koppen:
            idx = {k: j for j, k in enumerate(koppen) if k}
            code_j = next(j for k, j in idx.items() if k.startswith("opleidingscode") or k.startswith("crebo"))
            out = {}
            for r in rijen[i + 1:]:
                if not _is_code(r[code_j]):
                    continue
                def cel(k):
                    j = idx.get(k)
                    return r[j] if j is not None and j < len(r) else None
                out[_tekst(r[code_j])] = {
                    "dossiernaam": _tekst(cel("dossiernaam")),
                    "einde_instroom": _datum(cel("datumeindeinstroom")),
                    "einde_opleiding": _datum(cel("datumeindeopleiding")),
                    "vervangen_door": _tekst(cel("wordtvervangendoor")),
                }
            return out
    return {}


def parse_xlsx(content: bytes, editie: str, *, bron_url: str | None = None) -> CreboLijst:
    """Lees een SBB-CREBO-XLSX naar een genormaliseerde tabel (zonder omvangscontrole).

    Per kwalificatierij: ``crebo``, ``naam``, ``niveau``, ``prijsfactor``, ``soort_opleiding``,
    ``beroepsvereisten``, ``leerweg``, ``dossier_code``, ``dossier_naam``, ``domein``,
    ``geldig_van`` (begin van de editie), ``geldig_tot`` (einde opleiding van het dossier als dat in
    "Vervallen" staat), ``dossier_vervallen`` en ``bron`` (de ruwe cellen).
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("Installeer openpyxl: pip install 'riodata[sbb]'")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # zip/xml-fout in het bestand zelf
        raise CreboSchemaFout(f"Geen leesbaar XLSX-bestand: {e}") from e
    ws = _werkblad(wb, WERKBLAD)
    if ws is None:
        raise CreboSchemaFout(f"Werkblad 'Complete lijst' ontbreekt. Aanwezig: {[w.title for w in wb.worksheets]}")
    rijen = _rijen(ws)
    kop_i, pos = _kopregel(rijen)
    dossier = _dossierkolommen(rijen, kop_i)
    titel = _tekst(rijen[kop_i][0]) or _tekst(rijen[0][0]) or ""
    m = re.search(r"geldig vanaf\s+(\d{1,2}-\d{1,2}-\d{4})", titel, re.I)
    geldig_vanaf = _datum(m.group(1)) if m else None
    vervallen = _vervallen(wb)
    koppen = {j: _tekst(c) for j, c in enumerate(rijen[kop_i]) if _tekst(c) and j != 0}

    geldig_vanaf_sectie = geldig_vanaf
    uit = []
    domein = d_code = d_naam = None
    volgt_domein = False
    for rij in rijen[kop_i + 1:]:
        cel = lambda j: rij[j] if j is not None and j < len(rij) else None  # noqa: E731
        if not _is_code(cel(pos["crebo"])):
            titel_rij = _tekst(cel(0)) or ""
            tekst_naam = _tekst(cel(pos["naam"]))
            if "geldig vanaf" in titel_rij.lower():
                # Nieuwe sectie (bijv. Entree) met eigen kopregel en eigen ingangsdatum.
                m = re.search(r"geldig vanaf\s+(\d{1,2}-\d{1,2}-\d{4})", titel_rij, re.I)
                geldig_vanaf_sectie = _datum(m.group(1)) if m else geldig_vanaf_sectie
                domein = None
            elif tekst_naam and tekst_naam.lower() == "opleidingsdomein":
                volgt_domein = True
            elif tekst_naam and volgt_domein:
                domein = tekst_naam  # bijv. '1. Bouw en infra   79000'
                volgt_domein = False
            continue
        if dossier:
            naam = _tekst(cel(dossier["naam"]))
            code = _tekst(cel(dossier.get("code")))
            if code and _is_code(code):
                d_code, d_naam = code, naam
            elif naam and naam != d_naam:
                d_code, d_naam = None, naam
        ve = vervallen.get(d_code) if d_code else None
        uit.append({
            "crebo": _tekst(cel(pos["crebo"])),
            "naam": _tekst(cel(pos["naam"])),
            "niveau": _tekst(cel(pos["niveau"])),
            "prijsfactor": _getal(cel(pos.get("prijsfactor"))),
            "soort_opleiding": _tekst(cel(pos.get("soort_opleiding"))),
            "beroepsvereisten": _tekst(cel(pos.get("beroepsvereisten"))),
            "leerweg": _tekst(cel(pos.get("leerweg"))),
            "dossier_code": d_code,
            "dossier_naam": d_naam,
            "domein": re.sub(r"\s+", " ", domein) if domein else None,
            "geldig_van": geldig_vanaf_sectie,
            "geldig_tot": ve["einde_opleiding"] if ve else None,
            "dossier_vervallen": ve,
            "editie": editie,
            "bron": {koppen[j]: (cel(j).isoformat() if isinstance(cel(j), (dt.date, dt.datetime)) else cel(j))
                     for j in koppen},
        })
    return CreboLijst(
        editie=editie, rijen=tuple(uit), kolommen=tuple(koppen.values()),
        bron_url=bron_url, sha256=hashlib.sha256(content).hexdigest(),
        geldig_vanaf=geldig_vanaf, vervallen=vervallen,
    )


def valideer_download(content: bytes, content_type: str | None, lijst: CreboLijst) -> list[str]:
    """Controleer contenttype, bestandssignatuur, kolommen en omvang. Leeg = goedgekeurd."""
    problemen = []
    if content_type and content_type.split(";")[0].strip().lower() not in XLSX_CONTENTTYPES:
        problemen.append(f"contenttype '{content_type}' is geen XLSX")
    if not content.startswith(b"PK"):
        problemen.append("bestand is geen zip/XLSX (verkeerde signatuur)")
    if len(lijst.rijen) < MIN_RIJEN:
        problemen.append(f"slechts {len(lijst.rijen)} rijen (minimaal {MIN_RIJEN} verwacht)")
    if not any(r["crebo"] for r in lijst.rijen):
        problemen.append("geen enkele CREBO-code gevonden")
    return problemen


# ── editie-cache ─────────────────────────────────────────────────────────────

def cache_dir() -> Path:
    return Path(os.environ.get("RIODATA_CACHE", Path.home() / ".cache" / "riodata")) / "crebo"


def _manifest_pad() -> Path:
    return cache_dir() / "manifest.json"


def manifest() -> dict:
    """``{editie: {bron_url, sha256, bytes, rijen, gecontroleerd_op}}`` van gevalideerde edities."""
    try:
        return json.loads(_manifest_pad().read_text(encoding="utf-8"))
    except FileNotFoundError:
        return {}


def _bekende_edities() -> dict[str, str]:
    from . import sbb
    return {r["naam"]: r["url"] for r in sbb.resources("crebolijst")}


def _volgorde_sleutel(editie: str, meta: dict) -> tuple:
    """Sorteersleutel voor 'nieuwste editie': ingangsdatum uit het bestand, dan de volgorde in de catalogus.

    Bewust niet ``gecontroleerd_op``: de datum waarop wij een bestand controleren zegt niets over
    welke editie nieuwer is. Edities die uit de catalogus zijn verdwenen krijgen volgorde -1.
    """
    namen = list(_bekende_edities())
    return (meta.get("geldig_vanaf") or "", namen.index(editie) if editie in namen else -1)


def ververs(editie: str, *, client=None) -> CreboLijst:
    """Download een bekende editie, valideer die en zet haar in de cache.

    Een editie die de validatie niet haalt wordt niet in het manifest opgenomen en dus
    nooit ``latest``. Gooit ``CreboSchemaFout`` met de redenen.
    """
    import httpx
    bekend = _bekende_edities()
    if editie not in bekend:
        raise CreboEditieOnbekend(f"Onbekende editie '{editie}'. Bekend: {sorted(bekend)}")
    url = bekend[editie]
    r = (client or httpx).get(url, timeout=60, follow_redirects=True)
    r.raise_for_status()
    lijst = parse_xlsx(r.content, editie, bron_url=url)
    problemen = valideer_download(r.content, r.headers.get("content-type"), lijst)
    if problemen:
        raise CreboSchemaFout(f"Editie '{editie}' afgekeurd: " + "; ".join(problemen))
    return _bewaar(lijst, r.content)


def _bewaar(lijst: CreboLijst, content: bytes) -> CreboLijst:
    d = cache_dir()
    d.mkdir(parents=True, exist_ok=True)
    (d / f"{lijst.editie}.xlsx").write_bytes(content)
    nu = dt.datetime.now(dt.timezone.utc).isoformat(timespec="microseconds")
    m = manifest()
    m[lijst.editie] = {"bron_url": lijst.bron_url, "sha256": lijst.sha256, "bytes": len(content),
                       "rijen": len(lijst.rijen), "geldig_vanaf": lijst.geldig_vanaf, "gecontroleerd_op": nu}
    _manifest_pad().write_text(json.dumps(m, indent=2, ensure_ascii=False), encoding="utf-8")
    return CreboLijst(**{**lijst.__dict__, "gecontroleerd_op": nu})


def laad(editie: str = "latest") -> CreboLijst:
    """Laad een gecachete, gevalideerde editie.

    ``latest`` is de nieuwste gevalideerde editie: de hoogste ``geldig_vanaf`` (uit het bestand), bij
    gelijke datum de latere editie in de catalogus. Het opnieuw controleren van een oudere editie
    verandert dat niet. ``last_verified`` kiest juist de laatst gecontroleerde editie (dat is geen
    nieuwste editie). Is er geen gevalideerde editie, dan volgt een fout: er wordt niet teruggevallen
    op een ingebakken lijst. De checksum van het cachebestand wordt bij het laden opnieuw gecontroleerd.
    """
    m = manifest()
    if not m:
        raise CreboEditieOnbekend("Geen gevalideerde CREBO-editie in de cache. Draai crebo.ververs(<editie>) eerst.")
    if editie == "latest":
        editie = max(m, key=lambda e: _volgorde_sleutel(e, m[e]))
    elif editie == "last_verified":
        editie = max(m, key=lambda e: m[e]["gecontroleerd_op"])
    if editie not in m:
        raise CreboEditieOnbekend(f"Editie '{editie}' niet in cache. Beschikbaar: {sorted(m)}")
    content = (cache_dir() / f"{editie}.xlsx").read_bytes()
    if hashlib.sha256(content).hexdigest() != m[editie]["sha256"]:
        raise CreboSchemaFout(f"Cachebestand van '{editie}' komt niet overeen met de gecontroleerde checksum")
    lijst = parse_xlsx(content, editie, bron_url=m[editie]["bron_url"])
    return CreboLijst(**{**lijst.__dict__, "gecontroleerd_op": m[editie]["gecontroleerd_op"]})


# ── lookup ───────────────────────────────────────────────────────────────────

def zoek(code: str, editie: str | CreboLijst = "latest", peildatum: str | None = None) -> dict:
    """Zoek een CREBO-code deterministisch.

    Returns ``{"status", "code", "editie", "gecontroleerd_op", "matches"}`` met status:
    ``found`` (precies één geldige rij), ``not_found`` (geen) of ``ambiguous`` (meerdere
    geldige rijen: allemaal teruggegeven, nooit de eerste gekozen).

    ``peildatum`` (YYYY-MM-DD) filtert op geldigheid; een rij zonder begin- of einddatum telt als
    geldig aan die kant. Zonder peildatum tellen alle rijen met die code mee.
    """
    lijst = editie if isinstance(editie, CreboLijst) else laad(editie)
    sleutel = str(code).strip()
    matches = [r for r in lijst.rijen if r["crebo"] == sleutel]
    if peildatum:
        matches = [r for r in matches
                   if (r["geldig_van"] is None or r["geldig_van"] <= peildatum)
                   and (r["geldig_tot"] is None or r["geldig_tot"] >= peildatum)]
    status = "not_found" if not matches else "found" if len(matches) == 1 else "ambiguous"
    return {"status": status, "code": sleutel, "editie": lijst.editie, "geldig_vanaf": lijst.geldig_vanaf,
            "bron_url": lijst.bron_url, "sha256": lijst.sha256,
            "gecontroleerd_op": lijst.gecontroleerd_op, "peildatum": peildatum, "matches": matches}
